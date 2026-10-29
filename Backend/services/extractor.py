from io import BytesIO

from fastapi import UploadFile
from docx import Document

async def docx_text_extract(file: UploadFile):

    content = await file.read()

    document = Document(BytesIO(content))

    # paragraphs
    paragraphs = [
        {
            "text": p.text,
            "style": p.style.name
        }
        for p in document.paragraphs
        if p.text.strip()
    ]

    # tables
    tables = []

    for table in document.tables:
        rows = []

        for row in table.rows:
            rows.append([cell.text for cell in row.cells])

        tables.append(rows)

    # headers
    headers = [
        section.header.paragraphs[0].text
        for section in document.sections
        if section.header.paragraphs
    ]

    # footers
    footers = [
        section.footer.paragraphs[0].text
        for section in document.sections
        if section.footer.paragraphs
    ]

    # headings
    headings = [
        p.text
        for p in document.paragraphs
        if p.style.name.startswith("Heading") and p.text.strip()
    ]

    return {
        "filename": file.filename,
        "paragraphs": paragraphs,
        "headings": headings,
        "tables": tables,
        "headers": headers,
        "footers": footers
    }