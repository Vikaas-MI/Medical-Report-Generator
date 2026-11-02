from docx import Document
from pathlib import Path
import os
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


BASE_DIR = Path(__file__).resolve().parent.parent

def populate_template(data: dict):

    template_path = BASE_DIR / "uploads" / "template.docx"
    output_path = os.path.join("reports", "report.docx")

    document = Document(template_path)

    paragraphs = list(document.paragraphs)

    for paragraph in paragraphs:
        heading = paragraph.text.strip()

        if heading in data:
            content = data[heading]

            new_paragraph = OxmlElement("w:p")
            paragraph._p.addnext(new_paragraph)

            new_para = Paragraph(new_paragraph, paragraph._parent)
            new_para.add_run(content)

    document.save(output_path)

    return output_path